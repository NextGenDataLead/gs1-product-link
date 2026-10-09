"""The client's short issue report: failed products, issues by category, and why (2026-10-09)."""

from __future__ import annotations

from pathlib import Path

import openpyxl

from lib import issue_report_files
from lib.eligibility import NO_IMAGE, NO_TARGET, Eligibility, LinkIssue
from lib.issue_report import (
    LINK,
    MISSING_DATA,
    MISSING_VIDEO,
    NO_IMAGE_CATEGORY,
    NO_NAME,
    NO_VIDEO,
    NOT_IN_EXPORT,
    NOT_RUN,
    RUN_ERROR,
    TWO_VIDEOS,
    Issue,
    IssueReport,
    run_issues,
    selection_issues,
)
from lib.scope_report import HELD, IN_SCOPE, NOT_SELECTED, ScopeRow, UnitResult

A, B, C, D, E = (f"0871319500000{n}" for n in range(1, 6))


def _report(products: list[tuple[str, str]], issues: list[Issue]) -> IssueReport:
    return IssueReport(
        "Product issues - Acme", "Selection", tuple(g for g, _ in products), tuple(issues)
    )


def test_the_selection_report_says_what_the_data_screen_says() -> None:
    products = [(A, "Vergiet"), (B, "Deurmat"), (C, "Zaklamp"), (D, "Schuurspons"), (E, "Kam")]
    verdict = Eligibility(
        not_eligible={
            B: "missing data: dim_height (attr 3498), dim_width (attr 3520); two videos in nl",
            C: NO_IMAGE,
        },
        missing_video={D: "no confirmed video in fr"},
    )

    issues = selection_issues(products, exported={B, C, D, E}, verdict=verdict)

    assert [(i.gtin, i.category, i.reason) for i in issues] == [
        (A, NOT_IN_EXPORT, "the GS1 export has no row for this barcode"),
        (B, MISSING_DATA, "dim_height (attr 3498), dim_width (attr 3520)"),
        (C, NO_IMAGE_CATEGORY, NO_IMAGE),
        (B, TWO_VIDEOS, "two videos in nl"),
        (D, MISSING_VIDEO, "no confirmed video in fr"),
    ]
    report = _report(products, issues)
    assert report.failed == [(A, "Vergiet"), (B, "Deurmat"), (C, "Zaklamp")]
    assert report.with_warnings == 1, "a missing video is an issue, not a failure"
    assert [category for category, _ in report.by_category()] == [
        NOT_IN_EXPORT,
        MISSING_DATA,
        NO_IMAGE_CATEGORY,
        TWO_VIDEOS,
        MISSING_VIDEO,
    ]
    assert [i.category for i in report.reasons_of(B)] == [MISSING_DATA, TWO_VIDEOS]


def test_a_links_batch_reports_its_links_only() -> None:
    verdict = Eligibility(
        bad_link={
            A: LinkIssue("https://acme.test/st", "the page does not exist (404)"),
            B: LinkIssue(None, NO_TARGET),
        }
    )

    issues = selection_issues([(A, "x"), (B, "y"), (C, "z")], {A, B, C}, verdict)

    assert [(i.gtin, i.category, i.reason) for i in issues] == [
        (A, LINK, "the page does not exist (404) (https://acme.test/st)"),
        (B, LINK, NO_TARGET),
    ]


def test_a_video_hold_with_nothing_but_the_video_is_no_confirmed_video() -> None:
    issues = selection_issues(
        [(A, "x")], {A}, Eligibility(not_eligible={A: "no confirmed video in nl, fr"})
    )

    assert [(i.category, i.reason) for i in issues] == [(NO_VIDEO, "no confirmed video in nl, fr")]


def _scope(gtin: str, in_scope: str, **units: UnitResult) -> ScopeRow:
    return ScopeRow(cells=[gtin], gtin=gtin, in_scope=in_scope, units=units)


def test_the_run_report_covers_only_what_the_run_was_given() -> None:
    ok = UnitResult("ok", "https://acme.test/p/", "")
    rows = [
        _scope(A, IN_SCOPE, nl=ok, fr=ok),
        _scope(
            B,
            IN_SCOPE,
            nl=UnitResult("error", "", "PUT /links target URL for language nl does not serve: x"),
            fr=UnitResult(HELD, "", "missing_product_name: missing product_name.fr"),
        ),
        _scope(C, IN_SCOPE, nl=UnitResult("not run", "", ""), fr=ok),
        _scope(D, NOT_SELECTED, nl=UnitResult("not run", "", "")),
        _scope(A, IN_SCOPE, nl=ok, fr=ok),  # the same barcode on a second row is one product
    ]

    selected, issues = run_issues(
        rows, {A: "Vergiet", B: "Deurmat"}, writes_pages=True, without_video={A: ["fr"]}
    )

    assert selected == [A, B, C], "an unticked row is not this run's"

    assert [(i.gtin, i.category, i.reason) for i in issues] == [
        (B, NO_NAME, "fr: missing product_name.fr"),
        (B, RUN_ERROR, "nl: the page did not load, so no GS1 link was written"),
        (C, NOT_RUN, "nl: the run stopped before this product, or it was not confirmed"),
        (A, MISSING_VIDEO, "no video in fr"),
    ]


def test_a_links_run_never_reports_a_missing_video() -> None:
    ok = UnitResult("ok", "", "")
    _, issues = run_issues(
        [_scope(A, IN_SCOPE, nl=ok)], {}, writes_pages=False, without_video={A: ["nl"]}
    )

    assert issues == []


def test_the_pdf_and_the_workbook_carry_the_same_three_parts(tmp_path: Path) -> None:
    products = [(A, "Vergiet — groot"), (B, "Brosse à vaisselle")]
    issues = selection_issues(
        products, {B}, Eligibility(missing_video={B: "no confirmed video in fr"})
    )
    report = _report(products, issues)

    pdf, xlsx = issue_report_files.write(report, tmp_path, stem="issues")

    assert pdf.read_bytes().startswith(b"%PDF"), "a dash in a product name must not break it"
    book = openpyxl.load_workbook(xlsx)
    assert book.sheetnames == ["Total", "By category", "Failed products"]
    total = [row[0] for row in book["Total"].iter_rows(min_row=2, values_only=True)]
    assert "1 of 2 product(s) failed." in total
    failed = list(book["Failed products"].iter_rows(min_row=2, values_only=True))
    assert failed == [
        (A, "Vergiet — groot", "Not in the GS1 export: the GS1 export has no row for this barcode")
    ]
    by_category = list(book["By category"].iter_rows(min_row=2, values_only=True))
    assert [(row[0], row[1]) for row in by_category] == [
        ("Not in the GS1 export", "yes"),
        ("Live without a video", "no"),
    ]


def test_a_hold_in_every_language_is_said_once() -> None:
    held = UnitResult(HELD, "", "missing_mandatory_field: missing mandatory source data: image")
    _, issues = run_issues([_scope(A, IN_SCOPE, nl=held, fr=held)], {}, writes_pages=True)

    assert [(i.category, i.reason) for i in issues] == [(MISSING_DATA, "image")]
