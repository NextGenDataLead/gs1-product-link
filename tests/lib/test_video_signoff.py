"""The client's sign-off sheet, read against the mapping — and every shape an EAN arrives in.

The sheet is the one external dependency this pilot has been waiting on, and it arrives as a
spreadsheet somebody typed in. So the barcode column is the risk: Excel will hand a 13-digit EAN
back as a number, as a whole-number float, with a leading zero only if it was stored as text, and as
``8.7132E+12`` if the column was ever too narrow. Those are not edge cases — they are what the file
looks like. The first test here is the one that matters: every one of those spellings has to land on
the same canonical GTIN-14, because that is what :func:`lib.media_video.fully_mapped_gtins` will
compare against forever.

The second thing pinned here is that **nothing overwrites a confirmed row**. A GTIN in the
mapping is client sign-off; a sheet that disagrees with one is reported and never applied, and
``edits`` — the only thing a caller can write — must not contain it.
"""

from __future__ import annotations

import io

from openpyxl import Workbook

from lib.media_video import SKIP, VideoMap, VideoMapEntry
from lib.video_signoff import (
    AMBIGUOUS,
    BLANK,
    CONFLICT,
    FILL,
    REJECTED,
    UNCHANGED,
    column_options,
    columns,
    is_header,
    looks_like_an_export,
    plan,
    read_sheet,
)
from lib.xlsx import Grid, read_grid

#: The pilot's own languages, and a barcode that is really on its feed.
_LANGUAGES = ("nl", "fr")
_GTIN13 = "8713195008486"
_GTIN14 = "08713195008486"
_EXPORTED = frozenset({_GTIN14})


def _map(gtin: str = "", *, file: str = "Bulbman.mpg", language: str = "nl") -> VideoMap:
    """A mapping with one row, carrying whatever the test needs it to already say."""
    return VideoMap(by_language={language: [VideoMapEntry(file=file, gtin=gtin)], "fr": []})


def _sheet(
    *rows: tuple[str, str, str], header: tuple[str, str, str] = ("language", "file", "gtin")
) -> Grid:
    """A sign-off sheet as the reader would hand it over: header plus rows of text."""
    return Grid(list(header), [list(row) for row in rows])


def _only(grid: Grid, vmap: VideoMap, *, exported: frozenset[str] = _EXPORTED) -> object:
    """The single row a one-row sheet decides to, so a test can name one outcome."""
    decided = plan(grid, vmap, exported=exported, languages=_LANGUAGES)
    assert len(decided.rows) == 1, decided.rows
    return decided.rows[0]


def test_every_spelling_of_one_barcode_lands_on_the_same_gtin() -> None:
    """The whole reason this module normalises rather than trusting the cell.

    Each of these is the same product. A 13-digit EAN typed plain, the same stored as text with its
    leading zero, the whole-number float Excel produces from a numeric cell, and the two ways people
    write a barcode down by hand.
    """
    spellings = [_GTIN13, _GTIN14, "8713195008486.0", "871-3195008486", "8713 195008486"]

    for given in spellings:
        row = _only(_sheet(("nl", "Bulbman.mpg", given)), _map())
        assert row.outcome == FILL, (given, row)
        assert row.gtin == _GTIN14, given


def test_scientific_notation_is_refused_with_the_fix_in_the_message() -> None:
    """The commonest way a barcode arrives broken, and the least obvious to whoever sent it.

    Stripping the ``E`` and padding what is left would produce a real-looking GTIN-14 for a
    product nobody chose, so this refuses — and says what to do, because "expected digits" gets
    the identical file sent back.
    """
    row = _only(_sheet(("nl", "Bulbman.mpg", "8.7132E+12")), _map())

    assert row.outcome == REJECTED
    assert "scientific notation" in row.detail
    assert "text" in row.detail


def test_too_many_digits_is_refused_rather_than_padded() -> None:
    """``zfill`` does not truncate, so this would read as "not in the export" — the wrong fault."""
    row = _only(_sheet(("nl", "Bulbman.mpg", "871319500848612345")), _map())

    assert row.outcome == REJECTED
    assert "18 digits" in row.detail


def test_a_cell_that_is_not_a_barcode_at_all_is_refused() -> None:
    """``canon_gtin`` strips every non-digit, so "see tab 2" would become a plausible GTIN-14."""
    row = _only(_sheet(("nl", "Bulbman.mpg", "see tab 2")), _map())

    assert row.outcome == REJECTED
    assert "not a barcode" in row.detail


def test_skip_is_an_answer_and_fills_like_a_gtin() -> None:
    """ "No product for this video" is a decision the client is entitled to make, not a gap."""
    row = _only(_sheet(("nl", "Bulbman.mpg", "Skip")), _map())

    assert row.outcome == FILL
    assert row.gtin == SKIP


def test_a_blank_cell_is_not_signed_off_and_not_an_error() -> None:
    """Most of a first-round sheet is blank. Calling that a problem would bury the real ones."""
    assert _only(_sheet(("nl", "Bulbman.mpg", "")), _map()).outcome == BLANK


def test_a_confirmed_row_is_never_overwritten() -> None:
    """The safety property. The sheet may be stale, and a GTIN in the mapping is client sign-off."""
    decided = plan(
        _sheet(("nl", "Bulbman.mpg", _GTIN13)),
        _map("07391905003191"),
        exported=_EXPORTED,
        languages=_LANGUAGES,
    )
    row = decided.rows[0]

    assert row.outcome == CONFLICT
    assert "07391905003191" in row.detail
    assert decided.edits == {}, "a conflict must not reach the only thing a caller can apply"


def test_the_same_answer_in_a_different_width_is_not_a_conflict() -> None:
    """The mapping is written with 13-digit GTINs and the sheet may carry 14, or the reverse.

    Compared as text this reads as a disagreement, and every already-signed-off row in the file
    would come back as a conflict — which would make the import useless on the second round.
    """
    row = _only(_sheet(("nl", "Bulbman.mpg", _GTIN14)), _map(_GTIN13))

    assert row.outcome == UNCHANGED


def test_a_gtin_the_export_does_not_carry_is_refused() -> None:
    """A transposed digit usually produces a barcode no product has. That is the catchable case."""
    row = _only(_sheet(("nl", "Bulbman.mpg", "8713195008468")), _map())

    assert row.outcome == REJECTED
    assert "not in the export" in row.detail


def test_with_no_export_read_the_export_check_is_skipped() -> None:
    """Otherwise an import attempted before ``parse_export`` rejects every row for the same wrong
    reason, and the operator is sent to MyGS1 over a file they have not parsed yet."""
    row = _only(_sheet(("nl", "Bulbman.mpg", _GTIN13)), _map(), exported=frozenset())

    assert row.outcome == FILL


def test_a_language_the_client_invented_is_refused() -> None:
    """Filling a block no run reads is worse than refusing: it looks like the work was done."""
    row = _only(_sheet(("de", "Bulbman.mpg", _GTIN13)), _map())

    assert row.outcome == REJECTED
    assert "not a configured language" in row.detail


def test_a_filename_the_mapping_does_not_have_is_refused_as_a_filename() -> None:
    """Checked before the barcode on purpose: a good GTIN against an unknown file is still nothing
    this import can do, and "not in the export" would send somebody to MyGS1 over a typo."""
    row = _only(_sheet(("nl", "Bulbmann.mpg", _GTIN13)), _map())

    assert row.outcome == REJECTED
    assert "no nl row for this file" in row.detail


def test_rejections_name_the_spreadsheet_row_the_reader_is_looking_at() -> None:
    """Line 1 is the header to whoever opens the file, so the first data row is line 2."""
    decided = plan(
        _sheet(("nl", "Bulbman.mpg", ""), ("nl", "Bulbman.mpg", "nope")),
        _map(),
        exported=_EXPORTED,
        languages=_LANGUAGES,
    )

    assert [row.line for row in decided.rows] == [2, 3]


def test_our_own_report_is_read_back_without_mistaking_a_hint_for_an_answer() -> None:
    """The client gets ``report_video_candidates``'s sheet, which carries ``candidate_1_gtin``
    beside ``gtin``. Matched loosely, the first *suggestion* would be read as the decision."""
    grid = Grid(
        [
            "language",
            "file",
            "normalized",
            "state",
            "gtin",
            "candidate_1_gtin",
            "candidate_1_score",
        ],
        [["nl", "Bulbman.mpg", "bulbman", "unset", _GTIN13, "07391905003191", "0.93"]],
    )

    where = columns(grid)

    assert where is not None
    assert where["gtin"] == 4
    assert plan(grid, _map(), exported=_EXPORTED, languages=_LANGUAGES).rows[0].gtin == _GTIN14


def test_the_barcode_column_is_found_under_an_earlier_reports_name_for_it() -> None:
    """``current_gtin`` is what an earlier ``report_video_candidates`` called the column, and the
    operator still holds sheets in that shape — the one in ``test-uploads`` carries 99 answers.

    Safe by construction, because in both schemas that single column is *both* the mapping's current
    value and where the client writes theirs: an untouched row decides UNCHANGED, a filled one FILL,
    and one disagreeing with sign-off CONFLICT, which is never applied.
    """
    grid = _sheet(("fr", "Bulbman.mpg", _GTIN13), header=("language", "file", "current_gtin"))

    assert columns(grid) == {"language": 0, "file": 1, "gtin": 2}
    assert is_header(["language", "file", "state", "current_gtin"])


def test_the_columns_are_found_under_the_names_a_client_would_use() -> None:
    """Nobody outside this project calls a barcode a GTIN."""
    grid = _sheet(("fr", "Bulbman.mpg", _GTIN13), header=("Language Code", "Filename", "EAN"))

    assert columns(grid) == {"language": 0, "file": 1, "gtin": 2}


def test_a_sheet_with_no_barcode_column_decides_nothing() -> None:
    """Rather than reading some other column as barcodes. The caller says so on screen."""
    grid = _sheet(("nl", "Bulbman.mpg", "x"), header=("language", "file", "comments"))

    assert columns(grid) is None
    assert plan(grid, _map(), exported=_EXPORTED, languages=_LANGUAGES).rows == ()


def test_the_header_row_needs_both_a_barcode_and_a_file_column() -> None:
    """A title row above the table often carries one word, and a data row can carry "file"."""
    assert is_header(["language", "file", "gtin"])
    assert not is_header(["Video sign-off — round 2"])
    assert not is_header(["gtin"])


def _workbook(*rows: tuple[str, str, object]) -> io.BytesIO:
    """A real .xlsx, shaped like the file a client actually sends back.

    Three things about it are not incidental, and each has broken a reader here before: the table
    is on the **second** sheet, it starts **below a title row**, and the barcode goes in as a
    **number**, which is what a spreadsheet does with 13 digits unless the column was set to text.
    """
    wb = Workbook()
    wb.active["A1"] = "a pivot nobody asked for"
    sheet = wb.create_sheet("Sign-off")
    sheet["A1"] = "Video sign-off — round 2"
    sheet.append([])
    sheet.append(["Language Code", "Filename", "EAN", "Notes"])
    for row in rows:
        sheet.append([*row, ""])
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream


def test_a_real_workbook_is_read_from_a_stream_and_the_table_is_found() -> None:
    """The upload is never written to disk, so the reader is handed bytes rather than a path.

    This is the whole path at once: a stream, a second sheet, a title row above the table, the
    client's own column names, and a barcode that arrived as a number.
    """
    grid = read_grid(_workbook(("nl", "Bulbman.mpg", int(_GTIN13))), header_row=is_header)

    assert grid is not None
    assert grid.header == ["Language Code", "Filename", "EAN", "Notes"]

    decided = plan(grid, _map(), exported=_EXPORTED, languages=_LANGUAGES)

    assert [(row.outcome, row.gtin) for row in decided.rows] == [(FILL, _GTIN14)]


def test_a_numeric_cell_and_a_text_cell_are_the_same_barcode_through_a_real_file() -> None:
    """Checked through a workbook, not only through the coercion, because the cell *type* is where
    this goes wrong: whether a leading zero survives is decided in Excel, not here."""
    grid = read_grid(
        _workbook(("nl", "Bulbman.mpg", int(_GTIN13)), ("fr", "Bulbman.mpg", _GTIN14)),
        header_row=is_header,
    )

    assert grid is not None
    vmap = VideoMap(
        by_language={
            "nl": [VideoMapEntry(file="Bulbman.mpg", gtin="")],
            "fr": [VideoMapEntry(file="Bulbman.mpg", gtin="")],
        }
    )
    decided = plan(grid, vmap, exported=_EXPORTED, languages=_LANGUAGES)

    assert {row.gtin for row in decided.rows} == {_GTIN14}


def test_a_sheet_whose_columns_are_called_nothing_we_know_still_arrives_on_screen() -> None:
    """The reason the column picker can exist at all.

    The accepted spellings are a guess about somebody else's spreadsheet, and the guess was wrong
    about the only real sheet there is. A file we cannot recognise must still come back with its
    headings listed, so the operator can say which column is which — being told to go and rename
    things in Excel to match a list the tool never shows is not a fix.
    """
    wb = Workbook()
    sheet = wb.active
    sheet.append(["kolom A", "het filmpje", "streepjescode"])
    sheet.append(["nl", "Bulbman.mpg", _GTIN13])
    stream = io.BytesIO()
    wb.save(stream)

    grid = read_sheet(stream.getvalue())

    assert grid is not None
    assert columns(grid) is None, "nothing here is recognisable — that is the premise"
    assert column_options(grid) == {0: "kolom A", 1: "het filmpje", 2: "streepjescode"}


def test_the_operators_own_column_choice_overrides_every_guess() -> None:
    """What the picker sends back. Three positions, and the names are not consulted again."""
    grid = Grid(["kolom A", "het filmpje", "streepjescode"], [["nl", "Bulbman.mpg", _GTIN13]])

    decided = plan(
        grid,
        _map(),
        exported=_EXPORTED,
        languages=_LANGUAGES,
        where={"language": 0, "file": 1, "gtin": 2},
    )

    assert [(row.outcome, row.gtin) for row in decided.rows] == [(FILL, _GTIN14)]


def test_a_column_with_no_heading_is_still_offered_by_position() -> None:
    """It may be the one wanted, and a sheet exported from a report often has one."""
    grid = Grid(["language", "", "gtin"], [["nl", "Bulbman.mpg", _GTIN13]])

    assert column_options(grid)[1] == "(column 2, no heading)"


def test_the_fallback_header_never_swallows_a_title_row() -> None:
    """A one-cell title above the table is not a header, however loose the fallback is."""
    wb = Workbook()
    sheet = wb.active
    sheet.append(["Video sign-off — round 2"])
    sheet.append([])
    sheet.append(["taal", "video", "code"])
    sheet.append(["nl", "Bulbman.mpg", _GTIN13])
    stream = io.BytesIO()
    wb.save(stream)

    grid = read_sheet(stream.getvalue())

    assert grid is not None
    assert grid.header == ["taal", "video", "code"]


def test_two_sheet_rows_claiming_one_barcode_are_both_held_back() -> None:
    """What the pilot's own sheet does: `Roll Light Summer.mpg` and `…Winter.mpg`, one GTIN.

    Two genuine videos of one product, which the mapping cannot express. Applying both attaches
    **neither** — ``resolve`` returns ``None`` for an ambiguous pair — to a product that still
    counts as mapped, so the page publishes with no video and nothing says why.
    """
    vmap = VideoMap(
        by_language={
            "nl": [
                VideoMapEntry(file="Roll Light Summer.mpg", gtin=""),
                VideoMapEntry(file="Roll Light Winter.mpg", gtin=""),
            ],
            "fr": [],
        }
    )
    grid = _sheet(
        ("nl", "Roll Light Summer.mpg", _GTIN13), ("nl", "Roll Light Winter.mpg", _GTIN13)
    )

    decided = plan(grid, vmap, exported=_EXPORTED, languages=_LANGUAGES)

    assert [row.outcome for row in decided.rows] == [AMBIGUOUS, AMBIGUOUS]
    assert decided.edits == {}, "neither may be applied — one of them has to be chosen by a person"
    assert "Roll Light Winter.mpg" in decided.rows[0].detail, "each names its rival"


def test_a_fill_colliding_with_an_existing_confirmed_file_is_held_back() -> None:
    """The collision can come from the mapping rather than from the sheet."""
    vmap = VideoMap(
        by_language={
            "nl": [
                VideoMapEntry(file="Super Trap.mp4", gtin=_GTIN13),
                VideoMapEntry(file="Super Trap.mpg", gtin=""),
            ],
            "fr": [],
        }
    )

    decided = plan(
        _sheet(("nl", "Super Trap.mpg", _GTIN13)), vmap, exported=_EXPORTED, languages=_LANGUAGES
    )

    assert decided.rows[0].outcome == AMBIGUOUS
    assert "Super Trap.mp4" in decided.rows[0].detail


def test_a_collision_is_found_across_barcode_widths() -> None:
    """The pilot's real case, and the one that was invisible: 13-digit in the file, 14 in the sheet.

    Compared as text these are two different GTINs and the pair looks fine. ``resolve``
    canonicalises, finds two files, and attaches nothing.
    """
    vmap = VideoMap(
        by_language={
            "nl": [
                VideoMapEntry(file="Super Trap.mp4", gtin=_GTIN13),
                VideoMapEntry(file="Super Trap.mpg", gtin=""),
            ],
            "fr": [],
        }
    )

    decided = plan(
        _sheet(("nl", "Super Trap.mpg", _GTIN14)), vmap, exported=_EXPORTED, languages=_LANGUAGES
    )

    assert decided.rows[0].outcome == AMBIGUOUS


def test_one_barcode_in_two_languages_is_the_normal_case_and_not_a_clash() -> None:
    """A product needs a video in *every* language, so this must stay the ordinary path."""
    vmap = VideoMap(
        by_language={
            "nl": [VideoMapEntry(file="Bulbman.mpg", gtin="")],
            "fr": [VideoMapEntry(file="Bulbman FR.mpg", gtin="")],
        }
    )
    grid = _sheet(("nl", "Bulbman.mpg", _GTIN13), ("fr", "Bulbman FR.mpg", _GTIN13))

    decided = plan(grid, vmap, exported=_EXPORTED, languages=_LANGUAGES)

    assert [row.outcome for row in decided.rows] == [FILL, FILL]
    assert len(decided.edits) == 2


def test_many_files_may_be_skip_without_clashing() -> None:
    """``skip`` is not a product, so it cannot collide with another ``skip``."""
    vmap = VideoMap(
        by_language={
            "nl": [
                VideoMapEntry(file="Trailer A.mpg", gtin=""),
                VideoMapEntry(file="Trailer B.mpg", gtin=""),
            ],
            "fr": [],
        }
    )
    grid = _sheet(("nl", "Trailer A.mpg", "skip"), ("nl", "Trailer B.mpg", "skip"))

    decided = plan(grid, vmap, exported=_EXPORTED, languages=_LANGUAGES)

    assert [row.outcome for row in decided.rows] == [FILL, FILL]
    assert len(decided.edits) == 2


def test_a_gs1_export_is_recognised_as_one_not_read_as_a_sheet() -> None:
    """The export's own opening headings, as lib.gdsn_layout writes them for the demo client."""
    export = Grid(
        header=["Gtin", "TargetMarketCountryCode", "InformationProviderOfTradeItem"], rows=[]
    )
    sheet = Grid(header=["language", "file", "gtin"], rows=[])

    assert looks_like_an_export(export)
    assert not looks_like_an_export(sheet)  # a "gtin" column alone is an ordinary sign-off sheet
