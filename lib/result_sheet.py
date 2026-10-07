"""Assembling and writing the per-run result sheet.

:mod:`lib.scope_report` decides *what the rows say* and is pure — no filesystem, no config, no
clock. This is the other half: finding the documents one run consumed, joining them, and writing the
workbook. It lives in ``lib`` rather than in the script because **two entry points need it**:

* ``scripts/report_scope_result.py``, to rebuild a sheet for any run on demand;
* ``scripts/run_execute.py``, at the end of every run, so a publish produces its own record.

That second one is the reason this module exists. The sheet used to be built by the *shell*, in the
handler that had just finished streaming a run — which meant a publish driven from anywhere else
produced no sheet at all, and the moment it is most wanted is the moment a run has half-failed.

**Every default path is relative to the working directory**, like ``lib/run_files`` and
``lib/state``: the scripts run with the repository as their cwd.
"""

from __future__ import annotations

from pathlib import Path
from typing import TYPE_CHECKING, Any, Final, NamedTuple

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from lib.config import ProcessListConfig
from lib.eligibility import eligibility
from lib.input_layout import archive_path
from lib.process_list import ProcessListSheet, read_process_list
from lib.records import Plan, ProductRecord, RunOutcome
from lib.run_files import RESULT_NAME, SELECTION_NAME, UPLOAD_NAME, sibling
from lib.scope_report import build_rows, legend_grid, scope_grid, units_grid

if TYPE_CHECKING:
    from collections.abc import Callable

    from lib.config import ClientConfig

#: Widest a column is auto-sized to — an error string runs to a paragraph, and a sheet whose first
#: column is 300 characters wide is worse to work in than one that truncates on screen.
_MAX_COLUMN_WIDTH: Final = 60


class Sheets(NamedTuple):
    """The two lists this report joins, and where they were found.

    ``uploaded`` is the superset — the list as the operator sent it — and ``control`` is the ticked
    subset the run consumed. ``archived`` says whether ``uploaded`` is genuinely the upload; when it
    is not, the two are the same sheet and no deselected row can be named. ``from_run`` says whether
    they came from the run's own directory, which is the difference between a report about *this*
    batch and one about whatever is in ``input/`` now.
    """

    uploaded: ProcessListSheet
    control: ProcessListSheet
    archived: bool
    from_run: bool


class Built(NamedTuple):
    """What the write produced, for a caller that wants to report on it."""

    out: Path
    rows: list[Any]
    unreadable: int
    sheets: Sheets
    plan_used: bool


def scope_sheets(cfg: ClientConfig, run_path: Path) -> Sheets:
    """The uploaded list and the ticked list, **from the run's own directory** where it has them.

    This is the whole point of ``lib/run_files``: a run copies in the selection it consumed and the
    upload that selection came from, because ``input/`` means "what the *next* run will use" and is
    overwritten by the next batch. A report that read ``input/`` afterwards described someone else's
    rows — and it did, for every run, because it looked for ``selections.source.xlsx``, a path
    nothing has ever written. So it silently took the fallback and reported the live files.

    Falls back to ``input/`` for a legacy run that kept no copies, so those still produce a report.
    """
    assert cfg.process_list is not None  # guarded by the caller
    column = cfg.process_list.gtin_column

    def read(path: Path) -> ProcessListSheet:
        return read_process_list(ProcessListConfig(path=str(path), gtin_column=column))

    used, sent = sibling(run_path, SELECTION_NAME), sibling(run_path, UPLOAD_NAME)
    if used.is_file():
        control = read(used)
        if sent.is_file():
            return Sheets(read(sent), control, True, True)
        return Sheets(control, control, False, True)

    # No copies in the run directory: a run from before they were kept. The live files are the best
    # available answer and may well be a later batch's, which is what ``from_run`` exists to say.
    control = read_process_list(cfg.process_list)
    archive = archive_path(control.path)
    if not archive.is_file():
        return Sheets(control, control, False, False)
    return Sheets(read(archive), control, True, False)


def load_outcomes(path: Path) -> tuple[list[RunOutcome], int]:
    """Read a run log, keeping the rows that parse and counting the ones that do not.

    A truncated final line is normal for a run killed mid-write, and discarding the whole file over
    it would throw away the record exactly when it matters most — which is also when this report
    is most likely to be asked for.
    """
    outcomes: list[RunOutcome] = []
    unreadable = 0
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        try:
            outcomes.append(RunOutcome.model_validate_json(line))
        except ValueError:
            unreadable += 1
    return outcomes, unreadable


def load_products(path: Path) -> list[ProductRecord]:
    """The parsed catalogue, for the set of GTINs the export actually carries."""
    import json  # noqa: PLC0415 — one caller, one line

    data = json.loads(path.read_text(encoding="utf-8"))
    return [ProductRecord.model_validate(item) for item in data]


def load_plan(path: Path) -> Plan | None:
    """The plan, for its holds. ``None`` when it is absent or will not validate."""
    try:
        return Plan.model_validate_json(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None


def plan_is_for(plan: Plan, outcomes: list[RunOutcome]) -> bool:
    """Whether this plan plausibly belongs to this run.

    ``plan.json`` is overwritten by every ``run_plan``, so for anything but the newest run it is
    somebody else's document. A plan generated *after* the run it is being read beside describes
    work that run never saw, and its holds would be reported as that run's.

    An inequality and not an identity, which is the honest description of it: nothing links a
    plan to the run that consumed it. Recording that link is the same move as
    ``runs/{stamp}/inputs.json`` and would make this a fact rather than a heuristic.
    """
    return not outcomes or plan.generated_at <= max(outcome.ts for outcome in outcomes)


def read_plan(
    client_id: str,
    outcomes: list[RunOutcome],
    *,
    path: str | None = None,
    no_plan: bool = False,
    warn: Callable[[str], None] | None = None,
) -> Plan | None:
    """The plan whose holds this report names, or ``None`` — saying why through ``warn``."""
    if no_plan:
        return None
    resolved = Path(path) if path else Path("output") / client_id / "plan.json"
    plan = load_plan(resolved)
    if plan is None:
        if path and warn:
            warn(f"warning: {resolved} could not be read as a plan")
        return None
    if not plan_is_for(plan, outcomes):
        if warn:
            warn(
                f"warning: {resolved} was generated after this run, so it is a later run's plan — "
                f"its holds are not reported. Pass --plan to point at the right one, or --no-plan."
            )
        return None
    return plan


def build(  # noqa: PLR0913 — one keyword per CLI flag, which is the point of it
    cfg: ClientConfig,
    run_path: Path,
    *,
    listed: str | None = None,
    plan_path: str | None = None,
    no_plan: bool = False,
    products_path: str | None = None,
    out: str | None = None,
    warn: Callable[[str], None] | None = None,
) -> Built:
    """Write the result sheet for one run and return what went into it.

    Raises:
        OSError, ValueError, ProcessListError: If a document it needs cannot be read. Callers decide
            what that means — an exit code for the CLI, a warning for a run that has already
            published and must not be stopped by a report.
    """
    assert cfg.process_list is not None  # guarded by the caller
    outcomes, unreadable = load_outcomes(run_path)
    products = load_products(
        Path(products_path)
        if products_path
        else Path("output") / cfg.client_id / "data" / "products.json"
    )
    if listed:
        one = read_process_list(
            ProcessListConfig(path=listed, gtin_column=cfg.process_list.gtin_column)
        )
        sheets = Sheets(one, one, True, False)
    else:
        sheets = scope_sheets(cfg, run_path)

    plan = read_plan(cfg.client_id, outcomes, path=plan_path, no_plan=no_plan, warn=warn)
    rows = build_rows(
        sheets.uploaded,
        selected=sheets.control.listed_gtins(),
        exported={product.gtin14 for product in products},
        outcomes=outcomes,
        skipped=plan.skipped if plan else [],
        languages=cfg.wordpress.languages,
        not_eligible=_not_eligible(cfg, sheets.uploaded, products),
    )

    columns, grid = scope_grid(sheets.uploaded, rows, cfg.wordpress.languages)
    # Into the run's own directory, beside its log and the selection it consumed. It used to be
    # `{stamp}-scope.xlsx` in the runs folder — a fourth naming scheme for a document about one
    # run, sitting in a directory of other runs' documents.
    destination = Path(out) if out else sibling(run_path, RESULT_NAME)
    write_workbook(
        destination,
        [
            ("scope", columns, grid),
            ("units", *units_grid(outcomes, plan.skipped if plan else [])),
            ("legend", *legend_grid()),
        ],
    )
    return Built(destination, rows, unreadable, sheets, plan is not None)


def _not_eligible(
    cfg: ClientConfig, uploaded: ProcessListSheet, products: list[ProductRecord]
) -> dict[str, str]:
    """``{gtin14: why}`` for the upload's products the Data screen held — the same verdict.

    The batch file holds only the ticked rows, so the plan never sees a held one to explain it.
    Asked now rather than when the batch was chosen, which can differ only if the export or the
    video mapping changed in between — and then today's reason is the one worth reading. Empty
    when the holds cannot be decided: the rows then read ``not selected``, never a wrong reason.
    """
    named = uploaded.listed_gtins()
    verdict = eligibility(cfg, [product for product in products if product.gtin14 in named])
    return {} if verdict.problem else verdict.not_eligible


def write_workbook(path: Path, sheets: list[tuple[str, list[str], list[list[str]]]]) -> None:
    """Write the workbook, each sheet with a frozen, filterable header row."""
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for title, columns, rows in sheets:
        sheet = workbook.create_sheet(title)
        sheet.append(columns)
        for row in rows:
            sheet.append(row)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        _size_columns(sheet, columns, rows)
    workbook.save(path)
    workbook.close()


def _size_columns(sheet: Any, columns: list[str], rows: list[list[str]]) -> None:
    """Widen each column to its widest value, capped."""
    for index, name in enumerate(columns):
        longest = max((len(str(row[index])) for row in rows if index < len(row)), default=0)
        width = min(max(len(name), longest) + 2, _MAX_COLUMN_WIDTH)
        sheet.column_dimensions[get_column_letter(index + 1)].width = width
